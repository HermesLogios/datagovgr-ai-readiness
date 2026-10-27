r"""
criteria_views.py
Readable versions of the assessment criteria.

What it does:
  Reads the assessment criteria from criteria/criteria.csv, the source file, and writes two
  versions for reading:
     criteria/criteria.xlsx   formatted workbook: wrapped text, fixed header and first columns,
                              filters, core criteria shaded
     criteria/criteria.md     one section per criterion, for reading on GitHub
  Both files are generated: changes are made in criteria.csv and the script is run again.

Notes:
  - The script has no number because it is not a step of the analysis: it is run whenever
    criteria.csv changes. The function write_criteria_sheet is reused by the workbook of the
    manual assessment, so that both show the criteria in the same way.
  - The workbook is written without time stamps (fixed dates inside the file), so that running
    the script again on the same criteria.csv gives exactly the same file.

Usage (from the project root, with the .venv activated):
    python scripts\criteria_views.py

Code licence: EUPL-1.2
"""

from __future__ import annotations

import io
import math
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

ROOT = Path(__file__).resolve().parent.parent
CRITERIA_DIR = ROOT / "criteria"
SOURCE_FILE = CRITERIA_DIR / "criteria.csv"
XLSX_FILE = CRITERIA_DIR / "criteria.xlsx"
MD_FILE = CRITERIA_DIR / "criteria.md"

# Columns of criteria.csv.
SOURCE_COLUMNS = [
    "id", "level", "name", "question", "score_0", "score_1", "score_2", "not_applicable",
    "method", "how_to_check", "core", "applies_to", "typical_fix", "sources", "fair",
]

# Columns of the workbook: (column of criteria.csv, heading, width in Excel units).
SHEET_COLUMNS = [
    ("id", "ID", 6),
    ("level", "Level", 14),
    ("name", "Criterion", 22),
    ("core", "Core", 7),
    ("question", "Question", 34),
    ("score_0", "Score 0", 40),
    ("score_1", "Score 1", 40),
    ("score_2", "Score 2", 40),
    ("not_applicable", "Not applicable", 28),
    ("method", "Method", 16),
    ("how_to_check", "How to check", 48),
    ("applies_to", "Applies to", 11),
    ("typical_fix", "Typical fix", 28),
    ("sources", "Sources", 36),
    ("fair", "FAIR", 7),
]
FROZEN_CELL = "E2"          # the heading row and the columns ID to Core stay visible

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
CORE_FILL = PatternFill("solid", fgColor="FFF2CC")
LEVEL_BORDER = Border(top=Side(style="medium", color="808080"))
LINE_HEIGHT = 15            # points per line of Calibri 11
CHARS_PER_UNIT = 1.1        # characters per unit of column width (kept low, so rows are not cut)

# Fixed date written inside the workbook instead of the time of the run.
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_DOC_TIME = "2026-01-01T00:00:00Z"


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------

def relative(path: Path) -> str:
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)


def load_criteria(path: Path) -> list[dict]:
    """Reads criteria.csv and checks that it is complete."""
    if not path.exists():
        raise SystemExit(f"{relative(path)} not found.")
    frame = pd.read_csv(path, encoding="utf-8-sig", dtype=str, keep_default_na=False)
    missing = [c for c in SOURCE_COLUMNS if c not in frame.columns]
    if missing:
        raise SystemExit(f"{relative(path)} lacks the columns {missing}.")
    problems = []
    if frame["id"].duplicated().any():
        problems.append("repeated IDs: " + ", ".join(frame.loc[frame["id"].duplicated(), "id"]))
    for row in frame.to_dict("records"):
        for column in ["id", "level", "name", "question", "score_0", "score_1", "score_2", "core", "applies_to"]:
            if not row[column].strip():
                problems.append(f"{row['id'] or '(no ID)'}: empty {column}")
        if row["core"] not in ("yes", "no"):
            problems.append(f"{row['id']}: core must be yes or no")
    if problems:
        raise SystemExit(f"Errors in {relative(path)}:\n  " + "\n  ".join(problems))
    return frame.to_dict("records")


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def lines_needed(text: str, width: float) -> int:
    """Approximate number of lines that a text needs in a wrapped cell of the given width."""
    per_line = max(1, int(width * CHARS_PER_UNIT))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


def write_criteria_sheet(sheet, criteria: list[dict]) -> None:
    """Writes the criteria on a worksheet, one row per criterion, formatted for reading."""
    for index, (_, heading, width) in enumerate(SHEET_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=heading)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 2 * LINE_HEIGHT

    previous_level = None
    for row_number, criterion in enumerate(criteria, start=2):
        is_core = criterion["core"] == "yes"
        new_level = criterion["level"] != previous_level and previous_level is not None
        previous_level = criterion["level"]
        lines = 1
        for index, (column, _, width) in enumerate(SHEET_COLUMNS, start=1):
            value = criterion[column]
            cell = sheet.cell(row=row_number, column=index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if is_core:
                cell.fill = CORE_FILL
            if is_core and column in ("id", "name"):
                cell.font = Font(bold=True)
            if new_level:
                cell.border = LEVEL_BORDER
            lines = max(lines, lines_needed(value, width))
        sheet.row_dimensions[row_number].height = min(409, lines * LINE_HEIGHT + 4)

    sheet.freeze_panes = FROZEN_CELL
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def write_about_sheet(sheet, criteria: list[dict]) -> None:
    core = ", ".join(c["id"] for c in criteria if c["core"] == "yes")
    lines = [
        ("Assessment criteria", True),
        ("This workbook is a readable copy of criteria/criteria.csv, generated by "
         "scripts/criteria_views.py. Changes are made in the CSV file and the script is run again.", False),
        ("", False),
        ("Scale: 0 = not met, 1 = partly met, 2 = fully met. A score of n/a is allowed only where "
         "the column 'Not applicable' says so.", False),
        (f"Core criteria (rows shaded in yellow): {core}. They decide the group of each dataset: "
         "Ready when all core criteria score 2; Almost ready when every core criterion below 2 can "
         "be fixed with low effort; Not yet in all other cases.", False),
        ("Effort of a fix: low (metadata and documentation only), medium (a change in the file "
         "itself), high (a change in a system or in the source of the data).", False),
        ("Full rules: criteria/README.md.", False),
    ]
    sheet.column_dimensions["A"].width = 110
    for row_number, (text, bold) in enumerate(lines, start=1):
        cell = sheet.cell(row=row_number, column=1, value=text)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(bold=True, size=14) if bold else Font()
        sheet.row_dimensions[row_number].height = max(1, lines_needed(text, 110)) * LINE_HEIGHT + 4


def save_reproducible(workbook: Workbook, path: Path) -> None:
    """Saves the workbook with fixed dates, so that the same input always gives the same file."""
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)", rf"\g<1>{FIXED_DOC_TIME}\g<2>", text)
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(item.filename, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def write_xlsx(path: Path, criteria: list[dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Criteria"
    write_criteria_sheet(sheet, criteria)
    write_about_sheet(workbook.create_sheet("About"), criteria)
    save_reproducible(workbook, path)


# ---------------------------------------------------------------------------
# Markdown document
# ---------------------------------------------------------------------------

def md_text(text: str) -> str:
    """Text safe for Markdown: angle brackets would be read as HTML, a vertical bar as a table cell."""
    return str(text).replace("<", "&lt;").replace(">", "&gt;").replace("|", "/").strip()


def anchor(criterion: dict) -> str:
    """The link target that GitHub creates for the heading of a criterion."""
    heading = f"{criterion['id']} {criterion['name']}".lower()
    heading = re.sub(r"[^a-z0-9 _-]", "", heading)
    return heading.replace(" ", "-")


def build_markdown(criteria: list[dict]) -> str:
    core = ", ".join(c["id"] for c in criteria if c["core"] == "yes")
    lines = [
        "# Assessment criteria",
        "",
        "> Generated from `criteria.csv` by `scripts/criteria_views.py`. Changes are made in "
        "`criteria.csv`; this file is not edited by hand.",
        "",
        "Each criterion is scored 0 (not met), 1 (partly met) or 2 (fully met). "
        f"The core criteria ({core}) decide the group of each dataset; the rules are in "
        "[README.md](README.md).",
        "",
        "## Overview",
        "",
        "| ID | Criterion | Level | Core | Method |",
        "| --- | --- | --- | --- | --- |",
    ]
    for c in criteria:
        lines.append(f"| [{c['id']}](#{anchor(c)}) | {md_text(c['name'])} | {md_text(c['level'])} | "
                     f"{'yes' if c['core'] == 'yes' else ''} | {md_text(c['method'])} |")

    previous_level = None
    for c in criteria:
        if c["level"] != previous_level:
            lines += ["", f"## {c['level']}"]
            previous_level = c["level"]
        tags = [f"Level: {c['level']}"]
        if c["core"] == "yes":
            tags.append("core criterion")
        tags.append("applies to all datasets" if c["applies_to"] == "all"
                    else f"applies to {c['applies_to']} datasets only")
        lines += [
            "",
            f"### {c['id']} {c['name']}",
            "",
            f"*{' · '.join(tags)}*",
            "",
            f"**Question.** {md_text(c['question'])}",
            "",
            "| Score | Rule |",
            "| --- | --- |",
            f"| 0 | {md_text(c['score_0'])} |",
            f"| 1 | {md_text(c['score_1'])} |",
            f"| 2 | {md_text(c['score_2'])} |",
        ]
        if c["not_applicable"].strip():
            lines.append(f"| n/a | {md_text(c['not_applicable'])} |")
        lines += [
            "",
            f"**How to check** ({md_text(c['method'])}). {md_text(c['how_to_check'])}",
            "",
            f"**Typical fix.** {md_text(c['typical_fix'])}",
            "",
            f"**Sources.** {md_text(c['sources'])}" + (f" · **FAIR:** {md_text(c['fair'])}" if c["fair"].strip() else ""),
        ]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Main program
# ---------------------------------------------------------------------------

def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    criteria = load_criteria(SOURCE_FILE)
    write_xlsx(XLSX_FILE, criteria)
    MD_FILE.write_text(build_markdown(criteria), encoding="utf-8")

    levels = pd.Series([c["level"] for c in criteria]).value_counts(sort=False)
    print(f"Criteria read from {relative(SOURCE_FILE)}: {len(criteria)} "
          f"({', '.join(f'{level} {count}' for level, count in levels.items())})")
    print("Files:")
    print(f"  {relative(XLSX_FILE)}")
    print(f"  {relative(MD_FILE)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
